import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# -------------------------------------------------------------
# CONSTANTS & COLOR PALETTE (KARE Canteen Brand Theme)
# -------------------------------------------------------------
NAVY_DARK    = RGBColor(20, 28, 58)      # #141c3a
NAVY_PRIMARY = RGBColor(31, 42, 92)      # #1f2a5c
NAVY_LIGHT   = RGBColor(43, 58, 122)     # #2b3a7a
ORANGE       = RGBColor(252, 128, 25)    # #fc8019 (Swiggy accent)
GREEN        = RGBColor(36, 150, 63)     # #24963f (Pure veg / Security)
WHITE        = RGBColor(255, 255, 255)
BG_LIGHT     = RGBColor(247, 248, 250)   # #f7f8fa
CARD_BG      = RGBColor(255, 255, 255)
BORDER_GRAY  = RGBColor(233, 233, 235)   # #e9e9eb
TEXT_DARK    = RGBColor(28, 28, 40)      # #1c1c28
TEXT_MUTED   = RGBColor(104, 107, 120)   # #686b78
TEXT_LIGHT   = RGBColor(147, 149, 159)   # #93959f
CARD_ACCENT  = RGBColor(238, 242, 255)   # #eef2ff

FONT_TITLE = 'Segoe UI'
FONT_BODY  = 'Segoe UI'

def create_deck():
    prs = Presentation()
    prs.slide_width  = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Helper: Set background color of slide
    def set_slide_bg(slide, color):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = color
        bg.line.fill.background()
        return bg

    # Helper: Add standard slide header
    def add_header(slide, tag_text, title_text, subtitle_text=None, is_dark=False):
        # Category Tag Pill
        tag_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.45), Inches(2.8), Inches(0.35))
        tag_box.fill.solid()
        tag_box.fill.fore_color.rgb = ORANGE if is_dark else RGBColor(255, 244, 237)
        tag_box.line.fill.background()
        tf = tag_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_bottom = tf.margin_right = 0
        p = tf.paragraphs[0]
        p.text = tag_text.upper()
        p.alignment = PP_ALIGN.CENTER
        p.font.name = FONT_TITLE
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = WHITE if is_dark else ORANGE

        # Title
        tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.85), Inches(11.7), Inches(0.65))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_bottom = tf.margin_right = 0
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.name = FONT_TITLE
        p.font.size = Pt(24)
        p.font.bold = True
        p.font.color.rgb = WHITE if is_dark else NAVY_DARK

        # Subtitle
        if subtitle_text:
            p2 = tf.add_paragraph()
            p2.text = subtitle_text
            p2.font.name = FONT_BODY
            p2.font.size = Pt(12)
            p2.font.color.rgb = RGBColor(196, 201, 226) if is_dark else TEXT_MUTED

    # Helper: Add rounded card container
    def add_card(slide, x, y, w, h, bg_color=CARD_BG, border_color=BORDER_GRAY):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        if border_color:
            card.line.color.rgb = border_color
            card.line.width = Pt(1)
        else:
            card.line.fill.background()
        return card

    # =========================================================
    # SLIDE 1: TITLE SLIDE (Modern Dark Brand Theme)
    # =========================================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s1, NAVY_DARK)

    # Decorative header glow accent
    dec_bar = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.1), Inches(0.12), Inches(3.6))
    dec_bar.fill.solid()
    dec_bar.fill.fore_color.rgb = ORANGE
    dec_bar.line.fill.background()

    tb = s1.shapes.add_textbox(Inches(1.2), Inches(1.1), Inches(11.0), Inches(3.6))
    tf = tb.text_frame
    tf.word_wrap = True
    
    # Institution Badge
    p0 = tf.paragraphs[0]
    p0.text = "KALASALINGAM ACADEMY OF RESEARCH AND EDUCATION"
    p0.font.name = FONT_TITLE
    p0.font.size = Pt(13)
    p0.font.bold = True
    p0.font.color.rgb = ORANGE

    p0_sub = tf.add_paragraph()
    p0_sub.text = "Department of Computer Science & Engineering — Cyber Security"
    p0_sub.font.name = FONT_BODY
    p0_sub.font.size = Pt(12)
    p0_sub.font.color.rgb = RGBColor(180, 190, 220)
    p0_sub.space_after = Pt(20)

    p1 = tf.add_paragraph()
    p1.text = "SMART CANTEEN TOKEN SYSTEM"
    p1.font.name = FONT_TITLE
    p1.font.size = Pt(36)
    p1.font.bold = True
    p1.font.color.rgb = WHITE
    p1.space_after = Pt(8)

    p2 = tf.add_paragraph()
    p2.text = "Automated Web-Based Food Ordering, Real-Time Token Generation & Kitchen Display Platform"
    p2.font.name = FONT_BODY
    p2.font.size = Pt(16)
    p2.font.color.rgb = RGBColor(220, 225, 240)

    # 4 Meta Pill Badges at bottom
    badges = [
        ("Tech Stack", "Java 17 • Spring Boot 3 • JPA"),
        ("Database", "Supabase Cloud PostgreSQL"),
        ("Security", "Spring Security 6 • BCrypt • CSRF"),
        ("UI / UX", "Swiggy / Zomato Food Theme")
    ]
    card_w = 2.7
    gap = 0.3
    start_x = 0.8
    y_pos = 5.2

    for i, (b_title, b_val) in enumerate(badges):
        bx = start_x + i * (card_w + gap)
        c = add_card(s1, bx, y_pos, card_w, 1.4, bg_color=NAVY_PRIMARY, border_color=NAVY_LIGHT)
        tb_b = s1.shapes.add_textbox(Inches(bx + 0.15), Inches(y_pos + 0.15), Inches(card_w - 0.3), Inches(1.1))
        tf_b = tb_b.text_frame
        tf_b.word_wrap = True
        p_bt = tf_b.paragraphs[0]
        p_bt.text = b_title.upper()
        p_bt.font.name = FONT_TITLE
        p_bt.font.size = Pt(10)
        p_bt.font.bold = True
        p_bt.font.color.rgb = ORANGE
        p_bt.space_after = Pt(4)

        p_bv = tf_b.add_paragraph()
        p_bv.text = b_val
        p_bv.font.name = FONT_BODY
        p_bv.font.size = Pt(11)
        p_bv.font.color.rgb = WHITE

    # =========================================================
    # SLIDE 2: PROJECT OVERVIEW & BACKGROUND
    # =========================================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s2, BG_LIGHT)
    add_header(s2, "Executive Summary", "Project Overview & Campus Background", "Digital transformation of Kalasalingam Academy campus dining facilities")

    # 3 Column Cards
    col_w = 3.65
    c_gap = 0.38
    y_top = 1.7
    h_card = 5.0

    cards_s2 = [
        ("Campus Context", "Kalasalingam Academy Canteen", [
            ("Peak Hour Crunch", "Over 5,000+ students and faculty rush to campus canteens between 12:30 PM - 2:00 PM."),
            ("Physical Bottlenecks", "Manual counter queues lead to chaos, missed lectures, and delayed food preparation."),
            ("Digital Campus Mission", "KARE is driving smart campus digitization across student services & operations.")
        ], NAVY_PRIMARY),
        ("The Solution", "Smart Canteen Token System", [
            ("Direct Mobile Ordering", "Students order South Indian delicacies directly from their phones or laptops."),
            ("Unique Digital Tokens", "Automated collision-proof sequential tokens (T-0001, T-0002) generated instantly."),
            ("Kitchen Display System", "Kitchen chefs receive live order queues and manage prep states in one click.")
        ], ORANGE),
        ("Cyber Security Focus", "Enterprise Security Standards", [
            ("Role-Based Access (RBAC)", "Strict security partition between Students and Canteen Kitchen Admins."),
            ("BCrypt Password Hashing", "Zero plain-text credentials; salted one-way hashing for user safety."),
            ("Anti-CSRF & Parameterized JPA", "Full protection against Cross-Site Request Forgery and SQL Injection.")
        ], GREEN)
    ]

    for i, (badge, title, pts, accent_col) in enumerate(cards_s2):
        cx = 0.8 + i * (col_w + c_gap)
        add_card(s2, cx, y_top, col_w, h_card)

        # Top Accent Header Bar
        top_bar = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(cx + 0.3), Inches(y_top + 0.3), Inches(col_w - 0.6), Inches(0.4))
        top_bar.fill.solid()
        top_bar.fill.fore_color.rgb = accent_col
        top_bar.line.fill.background()
        tf_t = top_bar.text_frame
        p_t = tf_t.paragraphs[0]
        p_t.text = badge.upper()
        p_t.alignment = PP_ALIGN.CENTER
        p_t.font.name = FONT_TITLE
        p_t.font.size = Pt(11)
        p_t.font.bold = True
        p_t.font.color.rgb = WHITE

        # Content Text Box
        tb_c = s2.shapes.add_textbox(Inches(cx + 0.3), Inches(y_top + 0.9), Inches(col_w - 0.6), Inches(3.8))
        tf_c = tb_c.text_frame
        tf_c.word_wrap = True

        p_ct = tf_c.paragraphs[0]
        p_ct.text = title
        p_ct.font.name = FONT_TITLE
        p_ct.font.size = Pt(16)
        p_ct.font.bold = True
        p_ct.font.color.rgb = NAVY_DARK
        p_ct.space_after = Pt(14)

        for heading, desc in pts:
            p_h = tf_c.add_paragraph()
            p_h.text = f"•  {heading}: "
            p_h.font.name = FONT_BODY
            p_h.font.size = Pt(11)
            p_h.font.bold = True
            p_h.font.color.rgb = NAVY_PRIMARY

            p_d = tf_c.add_paragraph()
            p_d.text = f"   {desc}"
            p_d.font.name = FONT_BODY
            p_d.font.size = Pt(10.5)
            p_d.font.color.rgb = TEXT_MUTED
            p_d.space_after = Pt(8)

    # =========================================================
    # SLIDE 3: PROBLEM STATEMENT
    # =========================================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s3, BG_LIGHT)
    add_header(s3, "Challenges & Gaps", "Problem Statement: Traditional Canteen System", "Critical operational and security drawbacks observed in manual counter ticketing")

    # 2 Big Cards: Left = Traditional Pain Points, Right = Impact on Stakeholders
    add_card(s3, 0.8, 1.7, 5.7, 5.0)
    add_card(s3, 6.8, 1.7, 5.7, 5.0)

    # Left: Pain Points
    tb_l = s3.shapes.add_textbox(Inches(1.1), Inches(1.9), Inches(5.1), Inches(4.5))
    tf_l = tb_l.text_frame
    tf_l.word_wrap = True
    p = tf_l.paragraphs[0]
    p.text = "Operational Flaws in Manual Canteens"
    p.font.name = FONT_TITLE
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = RGBColor(190, 40, 40)
    p.space_after = Pt(14)

    flaws = [
        ("Chaotic Physical Queues", "Students crowd around billing counters for 15-25 minutes, leading to congestion and unhygienic conditions."),
        ("Token Collisions & Paper Loss", "Paper slips easily get lost, smudged, or mixed up between students, causing disputes at the serving counter."),
        ("No Stock Visibility", "Students wait in line only to discover their desired dish (e.g. Masala Dosa or Biryani) is sold out."),
        ("Accounting Inaccuracies", "Manual cash logging and paper bookkeeping create daily revenue leakage and reconciliation errors."),
        ("Security Vulnerabilities", "Paper coupons can be easily counterfeited or reused with zero audit logs or digital tracking.")
    ]
    for title, desc in flaws:
        ph = tf_l.add_paragraph()
        ph.text = f"❌  {title}"
        ph.font.name = FONT_TITLE
        ph.font.size = Pt(11.5)
        ph.font.bold = True
        ph.font.color.rgb = NAVY_DARK

        pd = tf_l.add_paragraph()
        pd.text = f"     {desc}"
        pd.font.name = FONT_BODY
        pd.font.size = Pt(10)
        pd.font.color.rgb = TEXT_MUTED
        pd.space_after = Pt(6)

    # Right: Impact Analysis
    tb_r = s3.shapes.add_textbox(Inches(7.1), Inches(1.9), Inches(5.1), Inches(4.5))
    tf_r = tb_r.text_frame
    tf_r.word_wrap = True
    p = tf_r.paragraphs[0]
    p.text = "Impact on College Stakeholders"
    p.font.name = FONT_TITLE
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = NAVY_PRIMARY
    p.space_after = Pt(14)

    impacts = [
        ("Students", "Significant study and recess break time lost standing in line; missed hot meal slots during tight 45-minute lunch intervals."),
        ("Canteen Chefs & Staff", "Kitchen staff overwhelmed by shouted orders and disorganized token sequences, causing erratic food delivery times."),
        ("Canteen Management", "Inability to predict demand, high food waste from over-preparation, and no analytics on peak-demand dishes."),
        ("Campus Security & Safety", "Overcrowding in enclosed dining halls violates campus health and safety protocols during peak college hours.")
    ]
    for title, desc in impacts:
        ph = tf_r.add_paragraph()
        ph.text = f"⚠️  {title}"
        ph.font.name = FONT_TITLE
        ph.font.size = Pt(11.5)
        ph.font.bold = True
        ph.font.color.rgb = ORANGE

        pd = tf_r.add_paragraph()
        pd.text = f"     {desc}"
        pd.font.name = FONT_BODY
        pd.font.size = Pt(10.5)
        pd.font.color.rgb = TEXT_MUTED
        pd.space_after = Pt(10)

    # =========================================================
    # SLIDE 4: OBJECTIVES & SCOPE
    # =========================================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s4, BG_LIGHT)
    add_header(s4, "Project Goals", "Project Objectives & Scope of Work", "Clear design targets and boundaries defined for the smart canteen implementation")

    # 4 Objective Cards in 2x2 Grid
    grid_w = 5.7
    grid_h = 2.35
    objs = [
        ("1. Contactless Self-Ordering", "Empower students to browse 12+ live South Indian food items with images, dynamic category filters, search capabilities, and local cart storage.", ORANGE),
        ("2. Atomic Sequential Token Generation", "Implement thread-safe synchronized token generation (T-0001, T-0002) in Spring Boot to guarantee zero duplication across simultaneous orders.", NAVY_PRIMARY),
        ("3. Real-Time Kitchen Display System (KDS)", "Provide kitchen staff with a live board that auto-refreshes every 5 seconds, enabling one-click order state transitions (PLACED -> READY -> COLLECTED).", GREEN),
        ("4. Real-Time Audio-Visual Student Alerts", "Deliver asynchronous order status polling with Web Audio API chime, floating toasts, and modal popups the moment food is marked READY.", NAVY_LIGHT)
    ]

    coords = [(0.8, 1.7), (6.8, 1.7), (0.8, 4.3), (6.8, 4.3)]
    for i, (title, desc, acc_col) in enumerate(objs):
        gx, gy = coords[i]
        add_card(s4, gx, gy, grid_w, grid_h)

        # Left color bar
        bar = s4.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(gx), Inches(gy), Inches(0.12), Inches(grid_h))
        bar.fill.solid()
        bar.fill.fore_color.rgb = acc_col
        bar.line.fill.background()

        tb_o = s4.shapes.add_textbox(Inches(gx + 0.3), Inches(gy + 0.25), Inches(grid_w - 0.5), Inches(grid_h - 0.5))
        tf_o = tb_o.text_frame
        tf_o.word_wrap = True
        po_t = tf_o.paragraphs[0]
        po_t.text = title
        po_t.font.name = FONT_TITLE
        po_t.font.size = Pt(14)
        po_t.font.bold = True
        po_t.font.color.rgb = NAVY_DARK
        po_t.space_after = Pt(8)

        po_d = tf_o.add_paragraph()
        po_d.text = desc
        po_d.font.name = FONT_BODY
        po_d.font.size = Pt(11)
        po_d.font.color.rgb = TEXT_MUTED

    # Scope Summary Footer Pill
    foot_pill = add_card(s4, 0.8, 6.8, 11.7, 0.45, bg_color=CARD_ACCENT, border_color=None)
    tb_fp = s4.shapes.add_textbox(Inches(0.8), Inches(6.8), Inches(11.7), Inches(0.45))
    tf_fp = tb_fp.text_frame
    p_fp = tf_fp.paragraphs[0]
    p_fp.text = "🎯 Scope: Web application + Desktop PWA + Supabase Cloud PostgreSQL database, ready for KARE Campus Canteen deployment."
    p_fp.alignment = PP_ALIGN.CENTER
    p_fp.font.name = FONT_TITLE
    p_fp.font.size = Pt(10.5)
    p_fp.font.bold = True
    p_fp.font.color.rgb = NAVY_PRIMARY

    # =========================================================
    # SLIDE 5: SYSTEM ARCHITECTURE & TECH STACK
    # =========================================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s5, BG_LIGHT)
    add_header(s5, "Technical Blueprint", "4-Tier System Architecture & Technology Stack", "Robust enterprise layering ensuring high throughput, modularity, and security")

    tiers = [
        ("Tier 1: Presentation", "Client Front-End", [
            "Thymeleaf 3 Server-Side Engine",
            "Swiggy / Zomato Modern CSS3",
            "Vanilla JavaScript (ES6+)",
            "HTML5 LocalStorage Cart",
            "Web Audio API Sound Alerts",
            "PWA Desktop Manifest"
        ], ORANGE),
        ("Tier 2: Application", "Core Services", [
            "Java 17 (LTS Runtime)",
            "Spring Boot 3.2.0 Framework",
            "Spring Security 6 (RBAC)",
            "Spring MVC Controllers",
            "RESTful JSON Polling APIs",
            "Embedded Tomcat 10 Server"
        ], NAVY_PRIMARY),
        ("Tier 3: Persistence", "Data Access Layer", [
            "Spring Data JPA",
            "Hibernate 6 ORM",
            "HikariCP Connection Pool",
            "Jakarta Bean Validation",
            "Transactional Services",
            "Sequential Token Generator"
        ], NAVY_LIGHT),
        ("Tier 4: Cloud DB", "Database Storage", [
            "Supabase Cloud PostgreSQL",
            "IPv4 Connection Pooler (6543)",
            "SSL/TLS Encrypted JDBC",
            "Automatic Schema Migration",
            "Cascading Foreign Keys",
            "High Availability 24/7"
        ], GREEN)
    ]

    tier_w = 2.7
    t_gap = 0.3
    for i, (t_name, t_sub, items, t_color) in enumerate(tiers):
        tx = 0.8 + i * (tier_w + t_gap)
        add_card(s5, tx, 1.7, tier_w, 5.0)

        # Header tag
        hbar = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(tx + 0.2), Inches(1.9), Inches(tier_w - 0.4), Inches(0.65))
        hbar.fill.solid()
        hbar.fill.fore_color.rgb = t_color
        hbar.line.fill.background()
        tf_h = hbar.text_frame
        p_h1 = tf_h.paragraphs[0]
        p_h1.text = t_name
        p_h1.alignment = PP_ALIGN.CENTER
        p_h1.font.name = FONT_TITLE
        p_h1.font.size = Pt(11)
        p_h1.font.bold = True
        p_h1.font.color.rgb = WHITE

        p_h2 = tf_h.add_paragraph()
        p_h2.text = t_sub
        p_h2.alignment = PP_ALIGN.CENTER
        p_h2.font.name = FONT_BODY
        p_h2.font.size = Pt(9)
        p_h2.font.color.rgb = RGBColor(230, 235, 250)

        # Items list
        tb_i = s5.shapes.add_textbox(Inches(tx + 0.2), Inches(2.7), Inches(tier_w - 0.4), Inches(3.8))
        tf_i = tb_i.text_frame
        tf_i.word_wrap = True
        for idx, item in enumerate(items):
            p = tf_i.paragraphs[0] if idx == 0 else tf_i.add_paragraph()
            p.text = f"✔  {item}"
            p.font.name = FONT_BODY
            p.font.size = Pt(10.5)
            p.font.color.rgb = NAVY_DARK
            p.space_after = Pt(8)

    # =========================================================
    # SLIDE 6: BLOCK DIAGRAM & DATA FLOW
    # =========================================================
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s6, BG_LIGHT)
    add_header(s6, "Architectural Diagram", "System Block Diagram & Component Interactions", "End-to-end data flow between Student Client, Web Server, Controllers, and Cloud Database")

    # Big Block Diagram Canvas
    canvas = add_card(s6, 0.8, 1.7, 11.7, 5.2, bg_color=WHITE, border_color=BORDER_GRAY)

    # Block 1: Clients (Left)
    c1 = add_card(s6, 1.1, 2.1, 2.3, 1.8, bg_color=CARD_ACCENT, border_color=NAVY_PRIMARY)
    tb_b1 = s6.shapes.add_textbox(Inches(1.2), Inches(2.2), Inches(2.1), Inches(1.6))
    tf_b1 = tb_b1.text_frame
    tf_b1.word_wrap = True
    p = tf_b1.paragraphs[0]
    p.text = "STUDENT CLIENT"
    p.font.name = FONT_TITLE
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = NAVY_DARK
    p.alignment = PP_ALIGN.CENTER
    for line in ["• Mobile / Laptop Browser", "• Swiggy UI & Cart (Local)", "• Polls /api/orders/{id}/status", "• Audio Chime on READY"]:
        p_l = tf_b1.add_paragraph()
        p_l.text = line
        p_l.font.size = Pt(8.5)
        p_l.font.color.rgb = TEXT_MUTED

    c2 = add_card(s6, 1.1, 4.4, 2.3, 1.8, bg_color=CARD_ACCENT, border_color=ORANGE)
    tb_b2 = s6.shapes.add_textbox(Inches(1.2), Inches(4.5), Inches(2.1), Inches(1.6))
    tf_b2 = tb_b2.text_frame
    tf_b2.word_wrap = True
    p = tf_b2.paragraphs[0]
    p.text = "KITCHEN ADMIN CLIENT"
    p.font.name = FONT_TITLE
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = ORANGE
    p.alignment = PP_ALIGN.CENTER
    for line in ["• Kitchen Display Screen (KDS)", "• 5-Sec Auto Refresh", "• Mark READY / COLLECTED", "• Menu CRUD & Inventory"]:
        p_l = tf_b2.add_paragraph()
        p_l.text = line
        p_l.font.size = Pt(8.5)
        p_l.font.color.rgb = TEXT_MUTED

    # Block 2: Security Filter Chain (Middle Left)
    s_box = add_card(s6, 3.8, 2.6, 2.3, 3.1, bg_color=NAVY_PRIMARY, border_color=None)
    tb_s = s6.shapes.add_textbox(Inches(3.9), Inches(2.7), Inches(2.1), Inches(2.9))
    tf_s = tb_s.text_frame
    tf_s.word_wrap = True
    p = tf_s.paragraphs[0]
    p.text = "SECURITY GATEWAY"
    p.font.name = FONT_TITLE
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = ORANGE
    p.alignment = PP_ALIGN.CENTER
    p.space_after = Pt(8)
    for line in [
        "Spring Security 6",
        "─────────────────",
        "• CSRF Token Validator",
        "• BCrypt Password Auth",
        "• Role Interception:",
        "  - /admin/** -> ADMIN",
        "  - /menu -> STUDENT",
        "• Session & Cookie Store"
    ]:
        p_l = tf_s.add_paragraph()
        p_l.text = line
        p_l.font.size = Pt(8.5)
        p_l.font.color.rgb = WHITE

    # Block 3: Controllers & Services (Middle Right)
    m_box = add_card(s6, 6.5, 2.1, 2.7, 4.1, bg_color=CARD_ACCENT, border_color=GREEN)
    tb_m = s6.shapes.add_textbox(Inches(6.6), Inches(2.2), Inches(2.5), Inches(3.9))
    tf_m = tb_m.text_frame
    tf_m.word_wrap = True
    p = tf_m.paragraphs[0]
    p.text = "SPRING BOOT CORE"
    p.font.name = FONT_TITLE
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = GREEN
    p.alignment = PP_ALIGN.CENTER
    p.space_after = Pt(6)
    for line in [
        "CONTROLLERS:",
        "• StudentController",
        "• AdminController",
        "• OrderApiController (REST)",
        "• AuthController",
        "──────────────────",
        "BUSINESS SERVICES:",
        "• OrderService (@Transactional)",
        "  - Token Generator (T-000X)",
        "  - Bill Computation",
        "• FoodItemService (Menu CRUD)",
        "• CustomUserDetailsService"
    ]:
        p_l = tf_m.add_paragraph()
        p_l.text = line
        p_l.font.size = Pt(8)
        p_l.font.color.rgb = NAVY_DARK

    # Block 4: Cloud Database Layer (Right)
    d_box = add_card(s6, 9.6, 2.6, 2.5, 3.1, bg_color=NAVY_DARK, border_color=ORANGE)
    tb_d = s6.shapes.add_textbox(Inches(9.7), Inches(2.7), Inches(2.3), Inches(2.9))
    tf_d = tb_d.text_frame
    tf_d.word_wrap = True
    p = tf_d.paragraphs[0]
    p.text = "SUPABASE CLOUD DB"
    p.font.name = FONT_TITLE
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = ORANGE
    p.alignment = PP_ALIGN.CENTER
    p.space_after = Pt(6)
    for line in [
        "PostgreSQL 15 Managed",
        "──────────────────",
        "TABLE ENTITIES:",
        "• users (Accounts/Roles)",
        "• food_items (Catalog)",
        "• orders (State/Total)",
        "• order_items (Items/Qty)",
        "• tokens (Unique Numbers)",
        "──────────────────",
        "HikariCP Connection Pool",
        "SSL/TLS Encrypted Driver"
    ]:
        p_l = tf_d.add_paragraph()
        p_l.text = line
        p_l.font.size = Pt(8)
        p_l.font.color.rgb = WHITE

    # =========================================================
    # SLIDE 7: WORKFLOW ARCHITECTURE & ORDER LIFECYCLE
    # =========================================================
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s7, BG_LIGHT)
    add_header(s7, "Process Flow", "Order Lifecycle: 6-Stage Workflow Architecture", "Chronological end-to-end trace from student dish selection to kitchen fulfillment")

    steps = [
        ("Step 1: Auth & Browse", "Student logs in with college email. Explores South Indian menu categorized into Breakfast, Meals, Snacks, and Beverages.", ORANGE),
        ("Step 2: Smart Cart", "Student adjusts quantities with interactive steppers. Cart dynamically calculates subtotal, campus discount, and net payable amount.", NAVY_PRIMARY),
        ("Step 3: Secure Order Dispatch", "Browser sends JSON payload via AJAX with CSRF token. Server validates stock and begins an atomic database transaction.", NAVY_LIGHT),
        ("Step 4: Token Generation", "OrderService increments sequential counter, binds unique token (e.g. T-0001), persists status as PLACED, and renders confirmation receipt.", GREEN),
        ("Step 5: Kitchen Prep", "Order surfaces on the Kitchen Display Board. Chefs prepare dishes and click 'Mark Ready', flipping database state to READY.", NAVY_PRIMARY),
        ("Step 6: Alert & Pickup", "Student device polling detects READY status, sounds Web Audio beep, shows popup. Student shows token at Counter 1-4 for collection.", ORANGE)
    ]

    sw = 3.65
    sh = 2.3
    scoords = [
        (0.8, 1.7), (4.85, 1.7), (8.9, 1.7),
        (0.8, 4.4), (4.85, 4.4), (8.9, 4.4)
    ]

    for i, (title, desc, scol) in enumerate(steps):
        sx, sy = scoords[i]
        add_card(s7, sx, sy, sw, sh)

        # Number badge
        nb = s7.shapes.add_shape(MSO_SHAPE.OVAL, Inches(sx + 0.25), Inches(sy + 0.25), Inches(0.45), Inches(0.45))
        nb.fill.solid()
        nb.fill.fore_color.rgb = scol
        nb.line.fill.background()
        tf_nb = nb.text_frame
        p_nb = tf_nb.paragraphs[0]
        p_nb.text = str(i + 1)
        p_nb.alignment = PP_ALIGN.CENTER
        p_nb.font.name = FONT_TITLE
        p_nb.font.size = Pt(13)
        p_nb.font.bold = True
        p_nb.font.color.rgb = WHITE

        tb_s = s7.shapes.add_textbox(Inches(sx + 0.8), Inches(sy + 0.22), Inches(sw - 0.95), Inches(sh - 0.35))
        tf_s = tb_s.text_frame
        tf_s.word_wrap = True
        ps_t = tf_s.paragraphs[0]
        ps_t.text = title
        ps_t.font.name = FONT_TITLE
        ps_t.font.size = Pt(12.5)
        ps_t.font.bold = True
        ps_t.font.color.rgb = NAVY_DARK
        ps_t.space_after = Pt(6)

        ps_d = tf_s.add_paragraph()
        ps_d.text = desc
        ps_d.font.name = FONT_BODY
        ps_d.font.size = Pt(10)
        ps_d.font.color.rgb = TEXT_MUTED

    # =========================================================
    # SLIDE 8: CYBER SECURITY IMPLEMENTATION (Dept of CSE - Cyber Security)
    # =========================================================
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s8, BG_LIGHT)
    add_header(s8, "Cyber Security Defense", "Cyber Security & Defensive Architecture", "Specialized defense implementations evaluated for the Department of CSE - Cyber Security")

    sec_cards = [
        ("🔐 Cryptographic Password Hashing", "BCrypt with Random Salt", [
            "BCrypt password hashing algorithm with auto-generated salt.",
            "Work factor ensures high computational cost, mitigating brute-force and rainbow table dictionary attacks.",
            "Passwords are never stored or logged in plain text anywhere in memory or the database."
        ], GREEN),
        ("🛡️ CSRF Token Validation", "Synchronizer Token Pattern", [
            "Spring Security generates unique, cryptographically random CSRF tokens per session.",
            "Token is embedded in Thymeleaf meta tags and validated on all state-altering POST requests.",
            "Guarantees that malicious third-party sites cannot submit unauthorized food orders."
        ], NAVY_PRIMARY),
        ("🚦 Role-Based Access Control (RBAC)", "URL Authorization & Interception", [
            "Strict authorization checks dividing ROLE_STUDENT and ROLE_ADMIN privileges.",
            "/admin/** endpoints strictly restricted to verified canteen administrative staff.",
            "Role-based dynamic UI rendering: Thymeleaf extras conditionally hide admin actions from students."
        ], ORANGE),
        ("💉 SQL Injection (SQLi) Immunity", "Parameterized JPA Queries", [
            "100% Parameterized queries executed through Spring Data JPA and Hibernate ORM.",
            "Zero raw string concatenations in database lookup or update statements.",
            "Prepared statements guarantee user input is treated strictly as data, never executable SQL code."
        ], NAVY_LIGHT),
        ("👁️ Data Privacy & Horizontal Isolation", "Student Order Boundary Enforcement", [
            "Students can strictly query and view only orders associated with their own user ID.",
            "Prevents horizontal privilege escalation where an attacker attempts to snoop other tokens.",
            "Order status polling API (/api/orders/{id}/status) verifies ownership before returning state."
        ], GREEN)
    ]

    # Display in 2 columns: Left 3 cards, Right 2 cards
    col_w8 = 5.7
    for i, (stitle, ssub, bullets, scolor) in enumerate(sec_cards):
        if i < 3:
            cx = 0.8
            cy = 1.7 + i * 1.75
            cw, ch = col_w8, 1.6
        else:
            cx = 6.8
            cy = 1.7 + (i - 3) * 2.65
            cw, ch = col_w8, 2.5

        add_card(s8, cx, cy, cw, ch)

        # Left color line
        sbar = s8.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(cx), Inches(cy), Inches(0.1), Inches(ch))
        sbar.fill.solid()
        sbar.fill.fore_color.rgb = scolor
        sbar.line.fill.background()

        tb_sc = s8.shapes.add_textbox(Inches(cx + 0.25), Inches(cy + 0.15), Inches(cw - 0.4), Inches(ch - 0.3))
        tf_sc = tb_sc.text_frame
        tf_sc.word_wrap = True

        p_st = tf_sc.paragraphs[0]
        p_st.text = f"{stitle} ({ssub})"
        p_st.font.name = FONT_TITLE
        p_st.font.size = Pt(12)
        p_st.font.bold = True
        p_st.font.color.rgb = NAVY_DARK
        p_st.space_after = Pt(4)

        for b in bullets:
            pb = tf_sc.add_paragraph()
            pb.text = f"• {b}"
            pb.font.name = FONT_BODY
            pb.font.size = Pt(9.5)
            pb.font.color.rgb = TEXT_MUTED
            pb.space_after = Pt(2)

    # =========================================================
    # SLIDE 9: DATABASE SCHEMA & ENTITY MODEL
    # =========================================================
    s9 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s9, BG_LIGHT)
    add_header(s9, "Relational Design", "Relational Database Schema & Data Dictionary", "5 core relational entities mapped in Supabase Cloud PostgreSQL with foreign key integrity")

    # Table of 5 entities
    t_shape = s9.shapes.add_table(6, 4, Inches(0.8), Inches(1.8), Inches(11.7), Inches(4.8))
    table = t_shape.table

    table.columns[0].width = Inches(1.8)
    table.columns[1].width = Inches(4.4)
    table.columns[2].width = Inches(2.7)
    table.columns[3].width = Inches(2.8)

    headers_tbl = ["TABLE NAME", "ATTRIBUTES & DATA TYPES", "CONSTRAINTS & KEYS", "PURPOSE & RELATION"]
    for j, h_name in enumerate(headers_tbl):
        cell = table.cell(0, j)
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY_DARK
        p = cell.text_frame.paragraphs[0]
        p.text = h_name
        p.font.name = FONT_TITLE
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = WHITE
        p.alignment = PP_ALIGN.CENTER

    data_rows = [
        ("users", "id (BIGSERIAL), name (VARCHAR),\nemail (VARCHAR), password_hash (VARCHAR),\nrole (VARCHAR)", "PK: id\nUNIQUE: email\nDEFAULT: 'STUDENT'", "Stores student & canteen admin credentials with BCrypt hashes."),
        ("food_items", "id (BIGSERIAL), name (VARCHAR),\nprice (DOUBLE), category (VARCHAR),\navailable (BOOLEAN), image_url (VARCHAR)", "PK: id\nDEFAULT: available=TRUE", "South Indian menu catalog with pricing, stock toggle, and image URLs."),
        ("orders", "id (BIGSERIAL), user_id (BIGINT),\ntotal (DOUBLE), status (VARCHAR),\ncreated_at (TIMESTAMP)", "PK: id\nFK: user_id -> users(id)\nON DELETE CASCADE", "Tracks order lifecycle: PLACED -> READY -> COLLECTED, linked to student."),
        ("order_items", "id (BIGSERIAL), order_id (BIGINT),\nfood_item_id (BIGINT), quantity (INT),\nprice (DOUBLE)", "PK: id\nFK: order_id -> orders\nFK: food_item_id -> food_items", "Junction table storing itemized breakdown and quantities per order."),
        ("tokens", "id (BIGSERIAL), order_id (BIGINT),\ntoken_no (VARCHAR), created_at (TIMESTAMP)", "PK: id\nUNIQUE: token_no\nUNIQUE FK: order_id -> orders", "Sequential digital token numbers (e.g. T-0001) for physical food pickup.")
    ]

    for i, row in enumerate(data_rows):
        for j, val in enumerate(row):
            cell = table.cell(i + 1, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = WHITE if i % 2 == 0 else CARD_ACCENT
            p = cell.text_frame.paragraphs[0]
            p.text = val
            p.font.name = FONT_BODY
            p.font.size = Pt(9.5)
            p.font.color.rgb = NAVY_DARK if j == 0 else TEXT_MUTED
            if j == 0:
                p.font.bold = True

    # =========================================================
    # SLIDE 10: USER EXPERIENCE & DESIGN INNOVATION
    # =========================================================
    s10 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s10, BG_LIGHT)
    add_header(s10, "Design Innovation", "Modern Food-Tech UI/UX (Swiggy / Zomato Inspired)", "Consumer-grade intuitive web interface optimized for both desktop and mobile devices")

    features_s10 = [
        ("🍔 Swiggy-Style Add Stepper", "Seamless micro-interaction: clicking 'ADD +' transforms immediately into a '− 1 +' quantity stepper, synchronizing automatically with browser LocalStorage.", ORANGE),
        ("🛒 Sticky Floating Cart Bar", "A bottom floating bar pops up the moment an item is added, displaying total item count and live subtotal with a quick 'View Cart ➔' action.", NAVY_PRIMARY),
        ("🌱 FSSAI Pure-Veg Indicators", "Every South Indian dish features an authentic green square & dot FSSAI vegetarian indicator alongside high-resolution food photography.", GREEN),
        ("🔔 Real-Time Audio Chime Alert", "Synthesized Web Audio API 880Hz chime rings automatically on the student's device the exact second their order status changes to READY.", NAVY_LIGHT),
        ("📱 Progressive Web App (PWA)", "Includes manifest.json, standalone window launcher (Desktop-App.bat), and zero-browser-chrome mode for a native application feel.", ORANGE),
        ("⚡ 1-Click Quick Demo Login", "Convenient demo login chips on the login screen for instant evaluation of both Student and Kitchen Admin roles without typing.", GREEN)
    ]

    for i, (ftitle, fdesc, fcol) in enumerate(features_s10):
        fx = 0.8 if i % 2 == 0 else 6.8
        fy = 1.7 + (i // 2) * 1.75
        add_card(s10, fx, fy, 5.7, 1.6)

        cbar = s10.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(fx), Inches(fy), Inches(0.1), Inches(1.6))
        cbar.fill.solid()
        cbar.fill.fore_color.rgb = fcol
        cbar.line.fill.background()

        tb_f = s10.shapes.add_textbox(Inches(fx + 0.25), Inches(fy + 0.15), Inches(5.3), Inches(1.3))
        tf_f = tb_f.text_frame
        tf_f.word_wrap = True
        pf_t = tf_f.paragraphs[0]
        pf_t.text = ftitle
        pf_t.font.name = FONT_TITLE
        pf_t.font.size = Pt(13)
        pf_t.font.bold = True
        pf_t.font.color.rgb = NAVY_DARK
        pf_t.space_after = Pt(4)

        pf_d = tf_f.add_paragraph()
        pf_d.text = fdesc
        pf_d.font.name = FONT_BODY
        pf_d.font.size = Pt(10)
        pf_d.font.color.rgb = TEXT_MUTED

    # =========================================================
    # SLIDE 11: ADVANTAGES & KEY BENEFITS
    # =========================================================
    s11 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s11, BG_LIGHT)
    add_header(s11, "Value Proposition", "Key Advantages & Multi-Stakeholder Benefits", "Quantifiable time savings, operational clarity, and hygiene improvements across the institution")

    # 3 Category Columns
    ben_cols = [
        ("For Students", NAVY_PRIMARY, [
            ("Zero Physical Waiting", "Eliminates standing in 20-minute lines during peak lunch hours; students wait comfortably at study tables."),
            ("Live Token Tracking", "Real-time order state visibility removes the uncertainty of when food will be prepared."),
            ("Audible Bell Alert", "Audio beep guarantees students won't miss their token call amidst canteen ambient noise.")
        ]),
        ("For Kitchen & Staff", ORANGE, [
            ("Digital Kitchen Board (KDS)", "Chefs work from an organized, auto-updating queue instead of confusing paper slips."),
            ("One-Click Status Changes", "Instantly updates students with a single button press ('Mark Ready' / 'Mark Collected')."),
            ("Menu Stock Management", "Admin can immediately mark finished items as 'Unavailable' to prevent out-of-stock orders.")
        ]),
        ("For College Administration", GREEN, [
            ("Hygiene & Crowd Control", "Resolves campus cafeteria overcrowding, supporting health and fire safety guidelines."),
            ("Daily Financial Summary", "Automated daily metrics report total tokens processed and exact revenue generated in INR."),
            ("24/7 Cloud Architecture", "Cloud PostgreSQL database requires zero local database servers or complex IT infrastructure.")
        ])
    ]

    for i, (btitle, bcolor, bitems) in enumerate(ben_cols):
        bx = 0.8 + i * 3.95
        add_card(s11, bx, 1.7, 3.8, 5.0)

        # Header box
        hdr = s11.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(bx + 0.2), Inches(1.9), Inches(3.4), Inches(0.5))
        hdr.fill.solid()
        hdr.fill.fore_color.rgb = bcolor
        hdr.line.fill.background()
        tf_h = hdr.text_frame
        p = tf_h.paragraphs[0]
        p.text = btitle
        p.alignment = PP_ALIGN.CENTER
        p.font.name = FONT_TITLE
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = WHITE

        tb_b = s11.shapes.add_textbox(Inches(bx + 0.25), Inches(2.6), Inches(3.3), Inches(3.9))
        tf_b = tb_b.text_frame
        tf_b.word_wrap = True

        for j, (h, d) in enumerate(bitems):
            ph = tf_b.paragraphs[0] if j == 0 else tf_b.add_paragraph()
            ph.text = f"✔  {h}"
            ph.font.name = FONT_TITLE
            ph.font.size = Pt(11)
            ph.font.bold = True
            ph.font.color.rgb = NAVY_DARK

            pd = tf_b.add_paragraph()
            pd.text = d
            pd.font.name = FONT_BODY
            pd.font.size = Pt(9.5)
            pd.font.color.rgb = TEXT_MUTED
            pd.space_after = Pt(10)

    # =========================================================
    # SLIDE 12: LIMITATIONS & CHALLENGES
    # =========================================================
    s12 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s12, BG_LIGHT)
    add_header(s12, "Critical Evaluation", "System Limitations & Technical Challenges", "Objective analysis of current technical constraints and mitigation strategies")

    challenges = [
        ("1. Campus Wi-Fi & Network Dependency", "Current Constraint", "The live token polling system requires continuous network connectivity. If student device loses connection, status alerts may lag.", "Mitigation / Defense", "Offline cart state preserved in LocalStorage; polling resumes automatically once connection re-establishes.", ORANGE),
        ("2. HTTP Polling vs WebSockets", "Current Constraint", "Polling /api/orders/{id}/status every 5 seconds generates lightweight HTTP requests that scale with the number of active users.", "Mitigation / Defense", "Polling interval is tuned to 5s with lightweight JSON response (<100 bytes); transition to WebSocket (STOMP) planned for Phase 2.", NAVY_PRIMARY),
        ("3. Digital Literacy for Contract Staff", "Current Constraint", "Canteen kitchen helpers might initially find digital screens unfamiliar compared to manual paper tickets.", "Mitigation / Defense", "High-contrast UI designed with large green and navy action buttons requiring only a single tap to transition status.", GREEN),
        ("4. Payment Gateway Integration", "Current Constraint", "Phase 1 implements server-side token accounting without direct online UPI / card payment processing.", "Mitigation / Defense", "Architecture includes modular order_id and total fields ready to plug in Razorpay / PhonePe webhooks.", NAVY_LIGHT)
    ]

    for i, (ctitle, ctype, cdesc, mtype, mdesc, ccol) in enumerate(challenges):
        cy = 1.7 + i * 1.28
        add_card(s12, 0.8, cy, 11.7, 1.15)

        cbar = s12.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(cy), Inches(0.1), Inches(1.15))
        cbar.fill.solid()
        cbar.fill.fore_color.rgb = ccol
        cbar.line.fill.background()

        tb_c = s12.shapes.add_textbox(Inches(1.05), Inches(cy + 0.1), Inches(11.3), Inches(0.95))
        tf_c = tb_c.text_frame
        tf_c.word_wrap = True

        p_t = tf_c.paragraphs[0]
        p_t.text = ctitle
        p_t.font.name = FONT_TITLE
        p_t.font.size = Pt(12)
        p_t.font.bold = True
        p_t.font.color.rgb = NAVY_DARK

        p_d = tf_c.add_paragraph()
        p_d.text = f"•  {ctype}: {cdesc}"
        p_d.font.name = FONT_BODY
        p_d.font.size = Pt(9.5)
        p_d.font.color.rgb = RGBColor(180, 50, 50)

        p_m = tf_c.add_paragraph()
        p_m.text = f"•  {mtype}: {mdesc}"
        p_m.font.name = FONT_BODY
        p_m.font.size = Pt(9.5)
        p_m.font.color.rgb = GREEN

    # =========================================================
    # SLIDE 13: FUTURE ENHANCEMENTS & ROADMAP
    # =========================================================
    s13 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s13, BG_LIGHT)
    add_header(s13, "Development Roadmap", "Future Enhancements & Scalability Plan", "Strategic technology extensions for campus-wide deployment across Kalasalingam Academy")

    roadmap = [
        ("Phase 2: Payment Gateway", "UPI & College RFID Smart Card Integration", [
            "Integrate Razorpay / PhonePe UPI payment gateway for seamless in-app payments.",
            "Connect KARE Student RFID identity cards via NFC readers for campus wallet deductions.",
            "Eliminate all cash exchanges at canteen counters."
        ], ORANGE),
        ("Phase 3: WebSocket Push", "Bi-directional WebSocket & Push Notifications", [
            "Upgrade status polling engine to Spring WebSockets (STOMP protocol over SockJS).",
            "Zero network polling overhead; server pushes order status instantaneously.",
            "Integrate Web Push Notifications and WhatsApp token alert triggers."
        ], NAVY_PRIMARY),
        ("Phase 4: Counter QR Scanner", "Automated Token Verification Hardware", [
            "Generate dynamic encrypted QR codes for every token on confirmation screen.",
            "Install budget barcode/QR scanners at Canteen Counter 1-4 for hands-free token redemption.",
            "Eliminates manual staff clicking for order collection."
        ], GREEN),
        ("Phase 5: AI Kitchen Predictor", "Machine Learning Demand Forecasting", [
            "Analyze historical ordering data to forecast daily food consumption per dish.",
            "Provide canteen chef with automated morning preparation estimates.",
            "Minimizes campus food wastage by up to 35%."
        ], NAVY_LIGHT)
    ]

    r_w = 2.7
    r_gap = 0.3
    for i, (r_phase, r_title, r_items, r_col) in enumerate(roadmap):
        rx = 0.8 + i * (r_w + r_gap)
        add_card(s13, rx, 1.7, r_w, 5.0)

        # Header tag
        rh = s13.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(rx + 0.2), Inches(1.9), Inches(r_w - 0.4), Inches(0.65))
        rh.fill.solid()
        rh.fill.fore_color.rgb = r_col
        rh.line.fill.background()
        tf_rh = rh.text_frame
        p_rh1 = tf_rh.paragraphs[0]
        p_rh1.text = r_phase
        p_rh1.alignment = PP_ALIGN.CENTER
        p_rh1.font.name = FONT_TITLE
        p_rh1.font.size = Pt(10.5)
        p_rh1.font.bold = True
        p_rh1.font.color.rgb = WHITE

        p_rh2 = tf_rh.add_paragraph()
        p_rh2.text = r_title
        p_rh2.alignment = PP_ALIGN.CENTER
        p_rh2.font.name = FONT_BODY
        p_rh2.font.size = Pt(8.5)
        p_rh2.font.color.rgb = RGBColor(230, 235, 250)

        tb_r = s13.shapes.add_textbox(Inches(rx + 0.2), Inches(2.7), Inches(r_w - 0.4), Inches(3.8))
        tf_r = tb_r.text_frame
        tf_r.word_wrap = True
        for idx, item in enumerate(r_items):
            p = tf_r.paragraphs[0] if idx == 0 else tf_r.add_paragraph()
            p.text = f"•  {item}"
            p.font.name = FONT_BODY
            p.font.size = Pt(10)
            p.font.color.rgb = NAVY_DARK
            p.space_after = Pt(8)

    # =========================================================
    # SLIDE 14: CONCLUSION & PROJECT SUMMARY
    # =========================================================
    s14 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s14, BG_LIGHT)
    add_header(s14, "Project Wrap-Up", "Conclusion & Key Takeaways", "Successful delivery of a secure, production-grade smart canteen automation system")

    # Left: Summary Card (7.5 inches)
    add_card(s14, 0.8, 1.7, 7.5, 5.0)
    tb_c1 = s14.shapes.add_textbox(Inches(1.1), Inches(1.9), Inches(6.9), Inches(4.5))
    tf_c1 = tb_c1.text_frame
    tf_c1.word_wrap = True

    p = tf_c1.paragraphs[0]
    p.text = "Summary of Completed Deliverables"
    p.font.name = FONT_TITLE
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = NAVY_PRIMARY
    p.space_after = Pt(14)

    conclusions = [
        ("Fully Operational Full-Stack Platform", "Built complete layered Java Spring Boot 3 web application with 64 files, Thymeleaf templates, and Swiggy-inspired CSS."),
        ("Live Cloud Database Integration", "Successfully connected to Supabase Cloud PostgreSQL via SSL connection pooler; tested with 12 South Indian menu items."),
        ("Real-Time Order-Token Lifecycle", "End-to-end token generation (T-0001), 5-second polling, audio beep notifications, and kitchen display status transitions tested and verified."),
        ("Cyber Security Compliance", "Demonstrated robust defense mechanisms (BCrypt password hashing, CSRF synchronizer tokens, RBAC URL filters, and SQLi-immune JPA queries)."),
        ("Campus Deployment Readiness", "Includes 1-click desktop runner (Desktop-App.bat), Cloudflare public mobile sharing (share-public-link.bat), and Dockerfile for cloud hosting.")
    ]

    for title, desc in conclusions:
        ph = tf_c1.add_paragraph()
        ph.text = f"✔  {title}"
        ph.font.name = FONT_TITLE
        ph.font.size = Pt(11.5)
        ph.font.bold = True
        ph.font.color.rgb = NAVY_DARK

        pd = tf_c1.add_paragraph()
        pd.text = f"    {desc}"
        pd.font.name = FONT_BODY
        pd.font.size = Pt(10)
        pd.font.color.rgb = TEXT_MUTED
        pd.space_after = Pt(8)

    # Right: Metric Highlights Card (3.9 inches)
    add_card(s14, 8.6, 1.7, 3.9, 5.0, bg_color=NAVY_DARK, border_color=None)
    tb_c2 = s14.shapes.add_textbox(Inches(8.8), Inches(1.9), Inches(3.5), Inches(4.5))
    tf_c2 = tb_c2.text_frame
    tf_c2.word_wrap = True

    p = tf_c2.paragraphs[0]
    p.text = "System Metrics"
    p.font.name = FONT_TITLE
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = ORANGE
    p.space_after = Pt(16)

    metrics = [
        ("75%", "Queue Time Reduction", "Estimated drop in student counter congestion during peak hours."),
        ("100%", "Token Collision Free", "Atomic thread-safe sequential generator eliminates duplicate slips."),
        ("12+", "South Indian Delicacies", "Pre-loaded menu catalog with live stock toggles and high-res imagery."),
        ("24 / 7", "Cloud Availability", "Supabase PostgreSQL hosted in Mumbai datacenter with SSL.")
    ]

    for num, lbl, sub in metrics:
        pn = tf_c2.add_paragraph()
        pn.text = num
        pn.font.name = FONT_TITLE
        pn.font.size = Pt(22)
        pn.font.bold = True
        pn.font.color.rgb = WHITE

        pl = tf_c2.add_paragraph()
        pl.text = lbl.upper()
        pl.font.name = FONT_TITLE
        pl.font.size = Pt(9.5)
        pl.font.bold = True
        pl.font.color.rgb = ORANGE

        ps = tf_c2.add_paragraph()
        ps.text = sub
        ps.font.name = FONT_BODY
        ps.font.size = Pt(9)
        ps.font.color.rgb = RGBColor(180, 190, 220)
        ps.space_after = Pt(10)

    # =========================================================
    # SLIDE 15: Q&A / THANK YOU SLIDE (Dark Brand Theme)
    # =========================================================
    s15 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s15, NAVY_DARK)

    # Center card container
    add_card(s15, 2.0, 1.2, 9.333, 5.1, bg_color=NAVY_PRIMARY, border_color=NAVY_LIGHT)

    tb_q = s15.shapes.add_textbox(Inches(2.3), Inches(1.5), Inches(8.7), Inches(4.5))
    tf_q = tb_q.text_frame
    tf_q.word_wrap = True

    pq1 = tf_q.paragraphs[0]
    pq1.text = "THANK YOU!"
    pq1.alignment = PP_ALIGN.CENTER
    pq1.font.name = FONT_TITLE
    pq1.font.size = Pt(40)
    pq1.font.bold = True
    pq1.font.color.rgb = ORANGE
    pq1.space_after = Pt(6)

    pq2 = tf_q.add_paragraph()
    pq2.text = "Questions & Feedback Welcome"
    pq2.alignment = PP_ALIGN.CENTER
    pq2.font.name = FONT_TITLE
    pq2.font.size = Pt(20)
    pq2.font.color.rgb = WHITE
    pq2.space_after = Pt(24)

    pq3 = tf_q.add_paragraph()
    pq3.text = "SMART CANTEEN TOKEN SYSTEM"
    pq3.alignment = PP_ALIGN.CENTER
    pq3.font.name = FONT_TITLE
    pq3.font.size = Pt(14)
    pq3.font.bold = True
    pq3.font.color.rgb = RGBColor(255, 200, 150)

    pq4 = tf_q.add_paragraph()
    pq4.text = "Kalasalingam Academy of Research and Education\nDepartment of Computer Science and Engineering (Cyber Security)"
    pq4.alignment = PP_ALIGN.CENTER
    pq4.font.name = FONT_BODY
    pq4.font.size = Pt(12)
    pq4.font.color.rgb = RGBColor(200, 210, 235)
    pq4.space_after = Pt(20)

    pq5 = tf_q.add_paragraph()
    pq5.text = "Live Demo Ready on http://localhost:8080 & Supabase Cloud\nStudent Login: student@kare.edu | Admin Login: admin@kare.edu"
    pq5.alignment = PP_ALIGN.CENTER
    pq5.font.name = FONT_BODY
    pq5.font.size = Pt(11)
    pq5.font.color.rgb = GREEN

    # Output paths
    out_dir = r"c:\Users\saip4\Downloads\java mini project\smart-canteen"
    out_file1 = os.path.join(out_dir, "Smart_Canteen_Token_System_Presentation.pptx")
    out_file2 = os.path.join(r"c:\Users\saip4\Downloads\java mini project", "Smart_Canteen_Token_System_Presentation.pptx")

    prs.save(out_file1)
    prs.save(out_file2)
    print(f"Presentation successfully created and saved to:\n  {out_file1}\n  {out_file2}")

if __name__ == "__main__":
    create_deck()
